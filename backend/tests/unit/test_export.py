"""Unit tests for P8: text extraction from .docx/.pdf, .docx export and the Understanding
Report."""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.export.docx_writer import build_docx
from app.export.extract import ExtractError, extract, extract_docx, extract_pdf, reflow
from app.export.understanding_report import build_report
from app.schemas.export import OwnershipNote, UnderstandingReportRequest


# ---------------------------------------------------------------- fixtures
def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Solar Tracker Report", level=0)
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Solar panels   convert light into electricity.")
    doc.add_paragraph("")
    doc.add_paragraph("Tracking raises the output.", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "secret table text"
    doc.add_paragraph("Table 1: Measured output.", style="Caption")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pdf(pages: list[list[str]]) -> bytes:
    """A minimal PDF with one Helvetica text line per string (no PDF library needed)."""
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]
    kids = []
    for lines in pages:
        ops = ["BT", "/F1 11 Tf", "14 TL", "72 720 Td"]
        for ln in lines:
            safe = ln.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({safe}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        content_id = len(objects) + 1
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_id = len(objects) + 1
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
            b"/Resources << /Font << /F1 %d 0 R >> >> >>" % (content_id, 0)
        )
        kids.append(page_id)
    font_id = len(objects) + 1
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects = [o.replace(b"/F1 0 0 R", b"/F1 %d 0 R" % font_id) for o in objects]
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % k for k in kids),
        len(kids),
    )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


# ---------------------------------------------------------------- extract
def test_extract_docx_keeps_order_headings_captions_and_skips_tables():
    out = extract_docx(make_docx())
    lines = out.text.split("\n\n")
    assert lines[:3] == [
        "Solar Tracker Report",
        "Introduction",
        "Solar panels convert light into electricity.",
    ]
    assert "- Tracking raises the output." in lines
    assert "Table 1: Measured output." in lines
    assert "secret table text" not in out.text
    assert out.warnings and "table" in out.warnings[0]


def test_extract_pdf_reflows_lines_and_drops_page_furniture():
    pages = [
        [
            "Lab Report",
            "Introduction",
            "Solar panels convert light into",
            "electricity for the grid.",
            "1",
        ],
        [
            "Lab Report",
            "The tracker follows the sun and gives more out-",
            "put than a fixed panel.",
            "2",
        ],
        ["Lab Report", "Results", "The output rose by 27 percent.", "3"],
    ]
    out = extract_pdf(make_pdf(pages))
    assert "Solar panels convert light into electricity for the grid." in out.text
    assert "gives more output than a fixed panel." in out.text  # hyphen re-joined
    assert "\n\nResults\n\n" in out.text
    assert "Lab Report" not in out.text  # running header removed
    assert not any(line.strip().isdigit() for line in out.text.split("\n"))  # page numbers removed


def test_reflow_keeps_a_sentence_across_a_page_break():
    text = reflow(["The pump was switched by a", "\f", "relay module rated at 10 A."])
    assert text == "The pump was switched by a relay module rated at 10 A."


@pytest.mark.parametrize(
    ("name", "data", "message"),
    [
        ("notes.txt", b"hello", "Only .docx and .pdf"),
        ("broken.docx", b"not a zip", "could not be read"),
        ("broken.pdf", b"%PDF-1.4 garbage", "could not be read"),
        ("scan.pdf", make_pdf([[]]), "no text layer"),
    ],
)
def test_extract_errors(name, data, message):
    with pytest.raises(ExtractError, match=message):
        extract(name, data)


def test_extract_size_limit():
    with pytest.raises(ExtractError, match="limit is 10 MB"):
        extract("big.pdf", b"0" * (10 * 1024 * 1024 + 1))


# ---------------------------------------------------------------- .docx export
REPORT = """Solar Tracker Report

Introduction
Solar panels convert light into electricity.
They work best facing the sun.

2.1 Tracking principle
The tracker follows the sun.

Results
- Output rose by 27%.
- Losses fell.

Figure 1: Output of both panels.
"""


def read(data: bytes):  # type: ignore[no-untyped-def]
    return Document(io.BytesIO(data))


def test_docx_export_structure_and_styles():
    doc = read(build_docx(REPORT, "Solar Tracker Report", "A. Student"))
    paras = [(p.style.name, p.text) for p in doc.paragraphs]
    assert paras[0] == ("Title", "Solar Tracker Report")
    assert paras[1][0] == "Subtitle" and paras[1][1].startswith("A. Student · ")
    texts = [t for _, t in paras]
    assert texts.count("Solar Tracker Report") == 1  # title not repeated from the text
    assert ("Heading 1", "Introduction") in paras
    assert ("Heading 2", "2.1 Tracking principle") in paras
    assert ("Heading 1", "Results") in paras
    assert ("List Bullet", "Output rose by 27%.") in paras
    assert ("Caption", "Figure 1: Output of both panels.") in paras
    intro = next(p for p in doc.paragraphs if p.text.startswith("Solar panels"))
    assert "They work best" in intro.text  # the line break inside a paragraph is kept
    assert doc.core_properties.title == "Solar Tracker Report"
    assert doc.core_properties.author == "A. Student"


def test_docx_export_ownership_note():
    note = OwnershipNote(score=64.4, student_share=0.38, passed_at="2026-09-26")
    doc = read(build_docx("Intro text here.", "T", note=note))
    last = doc.paragraphs[-1].text
    assert "Ownership Score 64/100" in last and "38%" in last and "2026-09-26" in last
    assert "Ownership note" not in " ".join(
        p.text for p in read(build_docx("Intro.", "T")).paragraphs
    )


# ---------------------------------------------------------------- understanding report
def test_understanding_report_sections():
    data = UnderstandingReportRequest(
        title="Solar Tracker",
        author="A. Student",
        ownership_score=64,
        student_share=0.38,
        ownership_parts=[
            {
                "label": "Text you wrote",
                "weight": 0.45,
                "value": 0.38,
                "detail": "380 of 1000 characters",
            },
            {"label": "Generic spots", "weight": 0.15, "value": None, "detail": "none found"},
        ],
        gate=[{"label": "Quiz", "value": 80, "required": 70, "passed": True}],
        gate_passed=True,
        mastered=["relay module"],
        weak=["duty cycle"],
        glossary=[
            {"term": "PLC", "definition": "programmable logic controller", "source": "document"}
        ],
        viva=[
            {
                "question": "What is a PLC?",
                "difficulty": "Define",
                "answer": "A computer.",
                "score": 55,
                "feedback": "Partly there.",
            }
        ],
        quiz_total=80,
        quiz_missed=[
            {
                "prompt": "The _____ switches the pump.",
                "answer_key": "relay",
                "your_answer": "valve",
            }
        ],
        teachback_coverage=62.5,
        teachback_explanation="We built a tracker.",
        teachback_missed=["servo"],
    )
    doc = read(build_report(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    for heading in (
        "Ownership",
        "Understanding checks",
        "Concepts",
        "Glossary",
        "Viva practice",
        "Revision list",
    ):
        assert heading in [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Ownership Score 64/100" in text
    assert "Q1. What is a PLC?" in text and "Your answer: A computer." in text
    assert "relay (you wrote: valve)" in text
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "programmable logic controller" in cells and "n/a" in cells and "passed" in cells
    revision = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert "duty cycle" in revision and "servo" in revision


def test_empty_understanding_report_still_builds():
    doc = read(build_report(UnderstandingReportRequest(title="Empty")))
    assert "No viva session yet." in "\n".join(p.text for p in doc.paragraphs)


# ---------------------------------------------------------------- endpoints
def test_files_extract_endpoint(client):
    r = client.post(
        "/files/extract", files={"file": ("report.docx", make_docx(), "application/octet-stream")}
    )
    assert r.status_code == 200 and r.json()["words"] > 5
    bad = client.post("/files/extract", files={"file": ("notes.txt", b"hi", "text/plain")})
    assert bad.status_code == 415
    scan = client.post(
        "/files/extract", files={"file": ("scan.pdf", make_pdf([[]]), "application/pdf")}
    )
    assert scan.status_code == 422


def test_export_endpoints(client):
    r = client.post("/export/docx", json={"text": REPORT, "title": "Solar Tracker Report!"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="Solar-Tracker-Report.docx"' in r.headers["content-disposition"]
    assert read(r.content).paragraphs[0].text == "Solar Tracker Report!"
    rep = client.post("/export/report", json={"title": "Solar Tracker"})
    assert (
        rep.status_code == 200
        and "Solar-Tracker-understanding-report.docx" in rep.headers["content-disposition"]
    )
    assert (
        client.post("/export/docx", json={"text": "word " * 3001, "title": "x"}).status_code == 413
    )
