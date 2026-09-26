"""The student's document as a Word file (CLAUDE.md §8 /export/docx).

Headings found by the segmenter become Word headings (numbered "2.1 …" headings one level
down), figure/table captions use the Caption style, "- " lines become bullets, and the rest are
normal paragraphs. An optional short ownership note can be added at the end.
"""

from __future__ import annotations

import io
import re
from datetime import date

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.shared import Pt

from app.core.segment import segment
from app.schemas.export import OwnershipNote

CAPTION_RE = re.compile(r"^\s*(?:Fig(?:ure)?\.?|Table)\s*\d+\s*[.:–—-]", re.IGNORECASE)
SUBHEADING_RE = re.compile(r"^\s*\d+\.\d+")
BULLET_RE = re.compile(r"^\s*[-•*]\s+")
MARKDOWN_HEADING_RE = re.compile(r"^\s*#+\s*")


def new_document(title: str, author: str | None) -> DocxDocument:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    props = doc.core_properties
    props.title = title
    props.author = author or ""
    props.comments = "Exported from OwnIt"
    return doc


def to_bytes(doc: DocxDocument) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_text(doc: DocxDocument, text: str, style: str | None = None) -> None:
    """A paragraph; single line breaks inside it are kept as Word line breaks."""
    lines = text.split("\n")
    p = doc.add_paragraph(style=style)
    for k, line in enumerate(lines):
        run = p.add_run(line.strip())
        if k < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)


def build_docx(
    text: str, title: str, author: str | None = None, note: OwnershipNote | None = None
) -> bytes:
    doc = new_document(title, author)
    doc.add_heading(title, level=0)
    byline = " · ".join(x for x in (author, date.today().strftime("%d %B %Y")) if x)
    doc.add_paragraph(byline, style="Subtitle")

    seg = segment(text)
    paragraphs = seg.paragraphs
    # the first line is the title again: don't repeat it
    if paragraphs and paragraphs[0].text.strip().lower() == title.strip().lower():
        paragraphs = paragraphs[1:]
    for p in paragraphs:
        body = p.text.strip()
        if not body:
            continue
        if p.is_heading:
            heading = MARKDOWN_HEADING_RE.sub("", body)
            doc.add_heading(heading, level=2 if SUBHEADING_RE.match(heading) else 1)
        elif CAPTION_RE.match(body):
            _add_text(doc, body, "Caption")
        elif all(BULLET_RE.match(ln) for ln in body.split("\n")):
            for ln in body.split("\n"):
                doc.add_paragraph(BULLET_RE.sub("", ln).strip(), style="List Bullet")
        else:
            _add_text(doc, body)

    if note is not None:
        doc.add_paragraph()
        box = doc.add_paragraph(style="Intense Quote")
        passed = f" Understanding checks passed on {note.passed_at}." if note.passed_at else ""
        box.add_run(
            f"Ownership note: Ownership Score {round(note.score)}/100; "
            f"{round(note.student_share * 100)}% of the text written by the author.{passed} "
            "Prepared with OwnIt (no AI models)."
        )
    return to_bytes(doc)
